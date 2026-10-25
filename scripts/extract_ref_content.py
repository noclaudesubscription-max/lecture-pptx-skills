"""
Extract the TEXT CONTENT (not styling) of a senior-faculty reference .pptx, for
use as a secondary content reference: what topics/terminology/examples were
previously taught. This is NOT a design/style extractor -- see
scripts/extract_ref_style.py's docstring and reference_style/NOTE.md for why
style is intentionally no longer derived from reference decks.

Usage:
    .venv\\Scripts\\python scripts\\extract_ref_content.py input\\"Unit 1 - Intro to C (1).pptx"

Writes reference_content/<slug>_content.md -- slide-by-slide title + body text
+ table contents, nothing about color/font/layout.
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent


def slugify(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", name.lower()).strip("-")
    return s


def extract(pptx_path: Path) -> str:
    prs = Presentation(str(pptx_path))
    lines = [f"# Reference Content — {pptx_path.name}", "",
             "> Secondary content reference only: topics, terminology, examples "
             "previously taught. Verify and improve before reuse -- do not copy "
             "verbatim, and never treat this as a design/layout/density reference.",
             ""]
    for idx, slide in enumerate(prs.slides, start=1):
        lines.append(f"## Slide {idx}")
        title = None
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            pass
        if title:
            lines.append(f"**Title:** {title}")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                for para in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in para.runs).strip()
                    if txt:
                        lines.append(f"- {txt}")
            if shape.has_table:
                lines.append("Table:")
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    lines.append("  | " + " | ".join(cells) + " |")
        lines.append("")
    return "\n".join(lines)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: extract_ref_content.py <path-to-reference.pptx>")
        sys.exit(1)
    src = Path(sys.argv[1])
    out_dir = ROOT / "reference_content"
    out_dir.mkdir(exist_ok=True)
    out_path = out_dir / f"{slugify(src.stem)}_content.md"
    out_path.write_text(extract(src), encoding="utf-8")
    print(f"Wrote {out_path}")
